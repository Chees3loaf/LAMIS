"""Raw-transcript parsing helpers shared without importing Tk."""
from __future__ import annotations
import re
from pathlib import Path
from typing import Dict, List, Optional

AUTO_DETECT_NOKIA = "Auto Detect Nokia"
SALES_BOM_IMPORT = "Sales BoM"
SCRIPT_OPTIONS: Dict[str, str] = {
    AUTO_DETECT_NOKIA: "", "Nokia PSI": "scripts.Nokia_PSI",
    "Nokia PSS": "scripts.Nokia_1830", "Nokia SAR": "scripts.Nokia_SAR_Raw",
    "Nokia IXR": "scripts.Nokia_IXR_Raw", "Ciena 6500": "scripts.Ciena_6500",
    "Ciena RLS": "scripts.Ciena_RLS", SALES_BOM_IMPORT: "__sales_bom_import__",
}
FAMILY_BY_MODULE = {"scripts.Nokia_PSI": "psi", "scripts.Ciena_RLS": "rls"}


def split_raw_output_by_commands(raw_text: str, commands: List[str]) -> List[str]:
    lines = raw_text.splitlines(); prompt = re.compile(r"^[^\#]*#\s*"); positions: Dict[int, int] = {}
    for require_prompt in (True, False):
        for line_index, line in enumerate(lines):
            if require_prompt and not prompt.match(line): continue
            candidate = prompt.sub("", line).strip()
            for command_index, command in enumerate(commands):
                if command_index not in positions and candidate.startswith(command): positions[command_index] = line_index; break
    ordered = sorted(positions.items(), key=lambda value: value[1]); result = [""] * len(commands)
    for index, (command_index, start) in enumerate(ordered):
        end = ordered[index + 1][1] if index + 1 < len(ordered) else len(lines)
        result[command_index] = "\n".join(lines[start + 1:end])
    return result


def read_excel_sheets(path: str) -> Dict[str, str]:
    import openpyxl
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return {name: "\n".join(str(row[0]) if row and row[0] is not None else "" for row in workbook[name].iter_rows(values_only=True)) for name in workbook.sheetnames}
    finally: workbook.close()


def normalize_device_id(value: str) -> str:
    original = (value or "").strip()
    cleaned = re.sub(r"^\s*\d{1,2}[-_/]\d{1,2}[-_/]\d{2,4}\s*-\s*\d{1,2}[\.:]\d{2}(?:\s*[AP]M)?(?:\s*[A-Z]{2,5})?\s*-\s*", "", original, count=1, flags=re.IGNORECASE).strip()
    return cleaned or original or "Manual"


def normalize_device_map(devices: Dict[str, str]) -> Dict[str, str]:
    result = {}
    for name, text in devices.items():
        base = normalize_device_id(name); key = base; counter = 2
        while key in result: key = f"{base}_{counter:02d}"; counter += 1
        result[key] = text
    return result


def read_text_folder(path: str) -> Dict[str, str]:
    devices = {}
    for source in sorted(Path(path).rglob("*")):
        if not source.is_file() or source.suffix.lower() != ".txt": continue
        base = normalize_device_id(source.stem); key = base; counter = 2
        while key in devices: key = f"{base}_{counter:02d}"; counter += 1
        devices[key] = source.read_text(encoding="utf-8", errors="replace")
    return devices


def detect_nokia_raw_script(raw_text: str, device_id: str = "") -> Optional[str]:
    haystack = f"{device_id}\n{raw_text}"
    for pattern, module in (
        (r"(?<![0-9A-Za-z])(7250|ixr(?:-r6d?)?)(?![0-9A-Za-z])", "scripts.Nokia_IXR_Raw"),
        (r"(?<![0-9A-Za-z])(7705|sar(?:-8)?)(?![0-9A-Za-z])", "scripts.Nokia_SAR_Raw"),
        (r"(?<![0-9A-Za-z])(1830|nokia\s*1830)(?![0-9A-Za-z])", "scripts.Nokia_1830"),
        (r"(?<![0-9A-Za-z])(psi|nokia(?:-[48]l)?)(?![0-9A-Za-z])", "scripts.Nokia_PSI"),
    ):
        if re.search(pattern, haystack, re.IGNORECASE): return module
    if re.search(r"^MDA\s+\d+/\d+\s+detail", raw_text, re.I | re.M) and re.search(r"^\s*Chassis\s+1\s+Detail", raw_text, re.I | re.M): return "scripts.Nokia_IXR_Raw"
    return None


def resolve_script_module(raw_text: str, device_id: str, script_name: str) -> str:
    if script_name == AUTO_DETECT_NOKIA:
        detected = detect_nokia_raw_script(raw_text, device_id)
        if not detected: raise RuntimeError(f"Could not auto-detect Nokia family for {device_id!r}. Select Nokia SAR or Nokia IXR explicitly.")
        return detected
    return SCRIPT_OPTIONS[script_name]
