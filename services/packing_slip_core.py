"""Toolkit-independent packing-slip source inspection and normalization."""
from __future__ import annotations

import logging
from pathlib import Path
from typing import Dict

import openpyxl
import pandas as pd

from utils.helpers import strip_dataframe_strings


def extract_workbook_metadata(path: str | Path) -> tuple[str, str, str, str]:
    customer = project = purchase_order = sales_order = ""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        summary_names = [name for name in workbook.sheetnames if "summary" in name.casefold()]
        report_names = []
        for name in workbook.sheetnames:
            marker = workbook[name]["B7"].value
            if isinstance(marker, str) and marker.strip().casefold().startswith("customer po"):
                report_names.append(name)
        if report_names:
            sheet = workbook[report_names[0]]
            values = [sheet[cell].value for cell in ("C5", "C6", "C7", "D7")]
            customer, project, purchase_order, sales_order = (
                "" if value is None or str(value).strip() in {"", "nan", "None"} else str(value).strip()
                for value in values
            )
        if not customer or not project:
            for name in workbook.sheetnames:
                if name in report_names or "summary" in name.casefold() or name.casefold() in {"bom", "inventory by site"}:
                    continue
                sheet = workbook[name]
                b5, b6 = sheet["B5"].value, sheet["B6"].value
                if (isinstance(b5, str) and b5.strip().casefold().startswith("customer") and
                        isinstance(b6, str) and b6.strip().casefold().startswith("project")):
                    if not customer and sheet["C5"].value not in (None, ""):
                        customer = str(sheet["C5"].value).strip()
                    if not project and sheet["C6"].value not in (None, ""):
                        project = str(sheet["C6"].value).strip()
                    break
        if (not customer or not project) and summary_names:
            sheet = workbook[summary_names[0]]
            if not customer and sheet["B7"].value not in (None, ""):
                customer = str(sheet["B7"].value).strip()
            if not project and sheet["D7"].value not in (None, ""):
                project = str(sheet["D7"].value).strip()
    finally:
        workbook.close()
    return customer, project, purchase_order, sales_order


def process_multisheet_device_file(path: str | Path) -> tuple[Dict[str, pd.DataFrame], dict[str, str], dict[str, str]]:
    processed: Dict[str, pd.DataFrame] = {}
    family_by_ip: dict[str, str] = {}
    display_ip_for_key: dict[str, str] = {}
    excel = pd.ExcelFile(path)
    for sheet_name in excel.sheet_names:
        raw = pd.read_excel(path, sheet_name=sheet_name, header=None)
        ip_address = sheet_name
        if raw.shape[0] > 4 and raw.shape[1] > 5:
            value = str(raw.iloc[4, 5]).strip()
            if value and value.casefold() != "nan":
                ip_address = value
        system_type = ""
        if raw.shape[0] > 6 and raw.shape[1] > 5:
            value = str(raw.iloc[6, 5]).strip()
            if value and value.casefold() != "nan":
                system_type = value
        bare_ip = ip_address
        if ip_address in processed:
            ip_address = f"{ip_address}_{sheet_name}"
        header_row = next((index for index, row in raw.iterrows() if "PART NUMBER" in " ".join(str(v).upper() for v in row if str(v).casefold() != "nan") or "SERIAL NUMBER" in " ".join(str(v).upper() for v in row if str(v).casefold() != "nan")), None)
        if header_row is None:
            logging.warning("Sheet %r: no packing-slip header row found; skipping", sheet_name)
            continue
        frame = pd.read_excel(path, sheet_name=sheet_name, header=header_row)
        strip_dataframe_strings(frame)
        frame = _truncate_additional_information(frame)
        relevant = [column for column in frame.columns if any(token in str(column).upper() for token in ("PART NUMBER", "SERIAL NUMBER", "DESCRIPTION"))]
        if relevant:
            frame = frame.dropna(subset=relevant, how="all")
            frame = frame[~frame[relevant].apply(lambda row: all(str(value).strip() in ("", "nan") for value in row), axis=1)]
        frame.insert(0, "System Name", sheet_name)
        if frame.empty:
            continue
        processed[ip_address] = frame.reset_index(drop=True)
        display_ip_for_key[ip_address] = bare_ip
        lowered = system_type.casefold()
        family_by_ip[ip_address] = "rls" if "rls" in lowered or "ciena" in lowered else "psi" if any(token in lowered for token in ("psi", "1830", "nokia")) else "default"
    return processed, family_by_ip, display_ip_for_key


def _truncate_additional_information(frame: pd.DataFrame) -> pd.DataFrame:
    sentinel = frame.apply(lambda row: row.astype(str).str.upper().str.contains(r"ADDITIONAL\s+NODE\s+INFORMATION", regex=True).any(), axis=1)
    if sentinel.any():
        cutoff = sentinel.idxmax()
        return frame.loc[:cutoff - 1] if cutoff > frame.index[0] else pd.DataFrame(columns=frame.columns)
    return frame


def process_packing_slip_frame(frame: pd.DataFrame) -> Dict[str, pd.DataFrame]:
    frame = _truncate_additional_information(frame)
    lowered = {column: str(column).casefold() for column in frame.columns}
    device_key = None
    for patterns in (("system name", "device", "ip address", " ip"), ("name",)):
        for pattern in patterns:
            device_key = next((column for column, label in lowered.items() if pattern in label or (label.strip() == "ip" and pattern == " ip")), None)
            if device_key is not None:
                break
        if device_key is not None:
            break
    if device_key is None:
        return {"Device_0": frame.reset_index(drop=True)}
    return {str(device_id): group.reset_index(drop=True) for device_id, group in frame.groupby(device_key, sort=False)}
