"""Tests for UI-independent Asset Import orchestration."""
from __future__ import annotations

from datetime import datetime

import openpyxl

from services.asset_import_service import run_asset_import


def test_run_asset_import_creates_backup_and_updates_inventory(tmp_path) -> None:
    inventory = tmp_path / "inventory.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "Summary"
    device = workbook.create_sheet("NODE-1")
    device["E15"] = "SERIAL-1"
    workbook.save(inventory)
    workbook.close()

    asset_doc = tmp_path / "assets.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["PO", "Serial #", "Asset"])
    sheet.append(["PO-42", "SERIAL-1", "ASSET-99"])
    workbook.save(asset_doc)
    workbook.close()

    messages: list[str] = []
    outcome = run_asset_import(
        inventory,
        asset_doc,
        progress=messages.append,
        now=datetime(2026, 8, 24, 12, 34, 56),
    )

    assert outcome.backup_path == tmp_path / "inventory.bak_20260824_123456.xlsx"
    assert outcome.backup_path.is_file()
    assert outcome.asset_records == 1
    assert outcome.result.rows_matched == 1
    assert any(message.startswith("Backup:") for message in messages)

    updated = openpyxl.load_workbook(inventory, data_only=False)
    assert updated["NODE-1"]["G15"].value == "ASSET-99"
    assert updated["NODE-1"]["C7"].value == "PO-42"
    updated.close()

    original = openpyxl.load_workbook(outcome.backup_path, data_only=False)
    assert original["NODE-1"]["G15"].value is None
    original.close()


def test_run_asset_import_rejects_missing_files(tmp_path) -> None:
    missing = tmp_path / "missing.xlsx"
    try:
        run_asset_import(missing, missing)
    except FileNotFoundError as exc:
        assert "Inventory workbook not found" in str(exc)
    else:
        raise AssertionError("missing input should fail")
