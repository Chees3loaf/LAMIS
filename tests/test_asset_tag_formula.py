"""Tests for the live Asset Tag XLOOKUP formula on device tabs.

Operator's request: rather than copying Asset Tag values from Summary
to device tabs as literals (the existing ``propagate_asset_tags_to_tabs``
behavior, which required a BoM rebuild to refresh), every device
tab's G15 should now carry a live XLOOKUP that pulls from Summary E
keyed on the device hostname. Edit Summary once, every device tab
follows automatically.
"""
from __future__ import annotations
import os
import shutil
import tempfile
import unittest
from unittest.mock import MagicMock

import openpyxl
import pandas as pd

from gui.workbook_builder import WorkbookBuilder


def _make_builder() -> WorkbookBuilder:
    db = MagicMock()
    db.db_path = ":memory:"
    db.lookup_part.return_value = ""
    return WorkbookBuilder(
        db_cache=db, template_path="", packing_slip_template="",
    )


# Excel stores post-2007 "future functions" under an ``_xlfn.`` prefix.
# Bare, Excel 365 renders the cell as ``=@XLOOKUP(...)`` and evaluates it
# to ``#NAME?``; prefixed, the formula bar shows the clean
# ``=XLOOKUP(F6,Summary!$D:$D,Summary!$E:$E,"",0,1)`` and it computes.
# The trailing ``&""`` keeps a listed-but-untagged device blank instead of
# showing the empty Summary cell as ``0``.
EXPECTED_FORMULA = '=_xlfn.XLOOKUP(F6,Summary!$D:$D,Summary!$E:$E,"",0,1)&""'
# Earlier generations that must be self-healed rather than preserved: the
# bare form leaves the column at #NAME?, the unprefixed-concat form leaves
# every untagged tab reading "0".
LEGACY_BARE_FORMULA = '=XLOOKUP(F6,Summary!$D:$D,Summary!$E:$E,"",0,1)'
LEGACY_NO_CONCAT_FORMULA = '=_xlfn.XLOOKUP(F6,Summary!$D:$D,Summary!$E:$E,"",0,1)'


class TestWriteAssetTagFormulaHelper(unittest.TestCase):
    """Unit tests for the new ``write_asset_tag_formula_to_chassis_row``
    helper — the workhorse called from every inventory builder."""

    def setUp(self):
        self.builder = _make_builder()
        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active

    def test_writes_formula_to_g15_by_default(self):
        wrote = self.builder.write_asset_tag_formula_to_chassis_row(self.ws)
        self.assertTrue(wrote)
        self.assertEqual(self.ws["G15"].value, EXPECTED_FORMULA)

    def test_formula_keys_on_hostname_in_f6(self):
        # F6 carries the device's system_name in every ATLAS-built
        # device tab — that's the lookup key the formula uses against
        # Summary's Device Name column.
        self.builder.write_asset_tag_formula_to_chassis_row(self.ws)
        formula = self.ws["G15"].value
        self.assertIn("F6", formula)
        self.assertIn("Summary!$D:$D", formula)
        self.assertIn("Summary!$E:$E", formula)

    def test_formula_carries_xlfn_prefix(self):
        # The whole point of the prefix: without it Excel treats XLOOKUP as
        # an unknown name, shows the legacy implicit-intersection "@" in the
        # formula bar, and the cell reads #NAME?.
        self.builder.write_asset_tag_formula_to_chassis_row(self.ws)
        formula = self.ws["G15"].value
        self.assertTrue(
            formula.startswith("=_xlfn.XLOOKUP("),
            f"Expected the _xlfn.-prefixed storage form, got {formula!r}",
        )
        self.assertNotIn("@", formula)

    def test_existing_formula_preserved_by_default(self):
        # Idempotent: a prior build already planted the current formula —
        # don't re-write it (avoids dirtying the workbook for no
        # reason on append-mode rebuilds).
        self.ws["G15"] = EXPECTED_FORMULA
        wrote = self.builder.write_asset_tag_formula_to_chassis_row(self.ws)
        self.assertFalse(wrote)

    def test_blank_tag_shows_empty_not_zero(self):
        # A device listed on Summary with no tag typed yet matches column D
        # and returns the empty column-E cell, which Excel renders as 0.
        # The trailing concat is what suppresses that.
        self.builder.write_asset_tag_formula_to_chassis_row(self.ws)
        self.assertTrue(
            self.ws["G15"].value.endswith('&""'),
            f'Expected a trailing &"" so untagged tabs read blank, '
            f'got {self.ws["G15"].value!r}',
        )

    def test_prefixed_formula_without_concat_is_upgraded(self):
        # The first pass at this fix resolved #NAME? but still displayed 0
        # for every untagged device; a rebuild must repair those too.
        self.ws["G15"] = LEGACY_NO_CONCAT_FORMULA
        wrote = self.builder.write_asset_tag_formula_to_chassis_row(self.ws)
        self.assertTrue(wrote)
        self.assertEqual(self.ws["G15"].value, EXPECTED_FORMULA)

    def test_legacy_bare_xlookup_is_upgraded_in_place(self):
        # Workbooks built before the prefix fix carry the bare form, which
        # Excel evaluates to #NAME?. A rebuild must repair it rather than
        # treat it as "already has a formula, leave alone".
        self.ws["G15"] = LEGACY_BARE_FORMULA
        wrote = self.builder.write_asset_tag_formula_to_chassis_row(self.ws)
        self.assertTrue(wrote)
        self.assertEqual(self.ws["G15"].value, EXPECTED_FORMULA)

    def test_unrelated_operator_formula_is_left_alone(self):
        # Self-healing must not extend to formulas the operator wrote.
        self.ws["G15"] = "=CONCAT(B15,\"-\",E15)"
        wrote = self.builder.write_asset_tag_formula_to_chassis_row(self.ws)
        self.assertFalse(wrote)
        self.assertEqual(self.ws["G15"].value, "=CONCAT(B15,\"-\",E15)")

    def test_key_cell_override_targets_packing_slip_layout(self):
        # Default packing-slip tabs keep the device name at C7, not F6.
        self.builder.write_asset_tag_formula_to_chassis_row(
            self.ws, key_cell="C7",
        )
        self.assertEqual(
            self.ws["G15"].value,
            '=_xlfn.XLOOKUP(C7,Summary!$D:$D,Summary!$E:$E,"",0,1)&""',
        )

    def test_key_cell_autodetected_from_layout(self):
        # F6 wins when populated (report layout); C7 is the fallback because
        # report tabs use C7 for the Customer PO string.
        self.ws["C7"] = "PO-123"
        self.ws["F6"] = "host-a"
        self.builder.write_asset_tag_formula_to_chassis_row(self.ws)
        self.assertIn("XLOOKUP(F6,", self.ws["G15"].value)

        ws2 = self.wb.create_sheet("pslip")
        ws2["C7"] = "host-b"
        self.builder.write_asset_tag_formula_to_chassis_row(ws2)
        self.assertIn("XLOOKUP(C7,", ws2["G15"].value)

    def test_literal_value_overwritten_by_formula(self):
        # An older workbook may have a literal asset tag at G15
        # (typed directly or stamped by the prior propagate flow).
        # New builds replace it with the formula so the live link
        # works going forward.
        self.ws["G15"] = "OLD_LITERAL_TAG_123"
        wrote = self.builder.write_asset_tag_formula_to_chassis_row(self.ws)
        self.assertTrue(wrote)
        self.assertEqual(self.ws["G15"].value, EXPECTED_FORMULA)

    def test_force_overwrites_existing_formula(self):
        self.ws["G15"] = '=SOME_OTHER_FORMULA()'
        wrote = self.builder.write_asset_tag_formula_to_chassis_row(
            self.ws, force=True,
        )
        self.assertTrue(wrote)
        self.assertEqual(self.ws["G15"].value, EXPECTED_FORMULA)


class TestPropagateLeavesFormulaCellAlone(unittest.TestCase):
    """``write_asset_tag_to_chassis_row`` is the legacy propagate
    primitive — it MUST skip cells that hold a formula now, or the
    propagate pass at BoM-build time would silently replace the live
    XLOOKUP with a frozen literal."""

    def setUp(self):
        self.builder = _make_builder()
        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active
        # Plant a chassis-shaped row at row 15 so the chassis-row
        # finder targets it.
        self.ws["B15"] = "Shelf 1"
        self.ws["C15"] = "Shelf"

    def test_skips_when_g15_is_formula(self):
        self.ws["G15"] = EXPECTED_FORMULA
        row = self.builder.write_asset_tag_to_chassis_row(self.ws, "NEW_TAG")
        self.assertEqual(row, 15)
        # Formula preserved.
        self.assertTrue(self.ws["G15"].value.startswith("="))

    def test_writes_literal_when_g15_was_empty(self):
        row = self.builder.write_asset_tag_to_chassis_row(self.ws, "TAG_42")
        self.assertEqual(row, 15)
        self.assertEqual(self.ws["G15"].value, "TAG_42")

    def test_overwrites_existing_literal(self):
        # Sanity: the formula-skip guard doesn't accidentally protect
        # legitimate operator-typed literals from being refreshed by
        # propagate.
        self.ws["G15"] = "STALE_TAG"
        self.builder.write_asset_tag_to_chassis_row(self.ws, "FRESH_TAG")
        self.assertEqual(self.ws["G15"].value, "FRESH_TAG")


class TestBuildReportWorkbookPlantsFormula(unittest.TestCase):
    """End-to-end: a fresh inventory report from the standard builder
    has every device tab's G15 holding the live XLOOKUP."""

    def setUp(self):
        self.tmp_dir = tempfile.mkdtemp(prefix="asset_tag_formula_")
        template_path = os.path.join(self.tmp_dir, "report_template.xlsx")
        tpl = openpyxl.Workbook()
        tpl.active.title = "Customer-project"
        tpl.create_sheet("Summary")
        # Device template carries G14="Asset Tag" header already in
        # the real install — replicate so the test sheet looks right.
        dev_tpl = tpl.create_sheet("DeviceTemplate")
        dev_tpl["G14"] = "Asset Tag"
        tpl.save(template_path)

        db = MagicMock()
        db.db_path = ":memory:"
        db.lookup_part.return_value = ""
        self.builder = WorkbookBuilder(
            db_cache=db, template_path=template_path,
            packing_slip_template="",
        )
        self.template_path = template_path

    def tearDown(self):
        shutil.rmtree(self.tmp_dir, ignore_errors=True)

    def _device_data(self, system_name: str) -> dict:
        df = pd.DataFrame([{
            "System Name": system_name,
            "System Type": "Test System",
            "Type": "Shelf",
            "Information Type": "Shelf",
            "Part Number": "P-1",
            "Serial Number": "S-1",
            "Description": "Test chassis",
            "Name": "Main Shelf",
            "Source": "LAN",
        }])
        return {"main": df}

    def test_g15_holds_formula_on_each_device_tab(self):
        out_path = os.path.join(self.tmp_dir, "out.xlsx")
        self.builder.build_report_workbook(
            {
                "10.0.0.1": self._device_data("device-one.example.com"),
                "10.0.0.2": self._device_data("device-two.example.com"),
            },
            out_path,
            customer="Test", project="Proj",
            customer_po="PO", sales_order="SO",
            append_mode=False,
        )
        # Reload preserving formulas.
        wb = openpyxl.load_workbook(out_path)
        device_tabs = [
            n for n in wb.sheetnames
            if n not in ("Summary", "Inventory by Site")
            and not n.startswith("DeviceTemplate")
            and not n.startswith("Customer")
        ]
        self.assertEqual(
            len(device_tabs), 2,
            f"Expected 2 device tabs, got {device_tabs!r}",
        )
        for tab in device_tabs:
            g15 = wb[tab]["G15"].value
            self.assertTrue(
                isinstance(g15, str) and g15.startswith("="),
                f"Tab {tab!r} G15 should be a formula, got {g15!r}",
            )
            # Hostname-keyed and pointing at Summary E.
            self.assertIn("F6", g15)
            self.assertIn("Summary", g15)
            # Stored in the form Excel actually resolves.
            self.assertEqual(g15, EXPECTED_FORMULA)


class TestPlantAcrossDeviceTabs(unittest.TestCase):
    """``plant_asset_tag_formula_on_device_tabs`` is the workbook-wide pass
    that runs before every save, so tabs an append-mode run didn't rewrite
    (and packing-slip tabs) are connected too."""

    def setUp(self):
        self.builder = _make_builder()
        self.wb = openpyxl.Workbook()
        self.wb.active.title = "Summary"
        self.wb["Summary"]["D9"] = "Device Name"
        self.wb["Summary"]["E9"] = "Asset Tag"

    def _device_tab(self, title: str, **cells) -> object:
        ws = self.wb.create_sheet(title)
        ws["F14"] = "DESCRIPTION"
        for addr, value in cells.items():
            ws[addr] = value
        return ws

    def test_plants_on_both_layouts_in_one_pass(self):
        report = self._device_tab("report-tab", C7="PO-9", F6="host-a")
        pslip = self._device_tab("pslip-tab", C7="host-b")
        planted = self.builder.plant_asset_tag_formula_on_device_tabs(self.wb)
        self.assertEqual(planted, 2)
        self.assertIn("XLOOKUP(F6,", report["G15"].value)
        self.assertIn("XLOOKUP(C7,", pslip["G15"].value)

    def test_skips_summary_and_aggregate_tabs(self):
        # An aggregate tab can carry text at C7 without being a device tab.
        for name in ("Inventory by Site", "BOM"):
            ws = self.wb.create_sheet(name)
            ws["F14"] = "DESCRIPTION"
            ws["C7"] = "Equipment Description"
        self.assertEqual(
            self.builder.plant_asset_tag_formula_on_device_tabs(self.wb), 0,
        )
        self.assertIsNone(self.wb["BOM"]["G15"].value)
        self.assertIsNone(self.wb["Summary"]["G15"].value)

    def test_skips_sheets_with_no_device_name(self):
        blank = self._device_tab("no-name-tab")
        self.assertEqual(
            self.builder.plant_asset_tag_formula_on_device_tabs(self.wb), 0,
        )
        self.assertIsNone(blank["G15"].value)

    def test_no_summary_sheet_means_no_plant(self):
        # The formula references Summary!$D:$D — without that sheet it could
        # only ever resolve to #REF!, so plant nothing.
        wb = openpyxl.Workbook()
        wb.active.title = "device-a"
        wb["device-a"]["F14"] = "DESCRIPTION"
        wb["device-a"]["F6"] = "host-a"
        self.assertEqual(
            self.builder.plant_asset_tag_formula_on_device_tabs(wb), 0,
        )
        self.assertIsNone(wb["device-a"]["G15"].value)

    def test_upgrades_legacy_formula_across_tabs(self):
        stale = self._device_tab("stale-tab", F6="host-a")
        stale["G15"] = LEGACY_BARE_FORMULA
        self.assertEqual(
            self.builder.plant_asset_tag_formula_on_device_tabs(self.wb), 1,
        )
        self.assertEqual(stale["G15"].value, EXPECTED_FORMULA)

    def test_adds_missing_asset_tag_header(self):
        # The Ciena RLS / Nokia PSI packing-slip templates ship without the
        # G14 label; the pass fills it so the column isn't unlabeled.
        tab = self._device_tab("rls-pslip", F6="host-a")
        self.assertIsNone(tab["G14"].value)
        self.builder.plant_asset_tag_formula_on_device_tabs(self.wb)
        self.assertEqual(tab["G14"].value, "Asset Tag")

    def test_existing_header_not_overwritten(self):
        tab = self._device_tab("has-header", F6="host-a")
        tab["G14"] = "Asset Tag"
        self.builder.plant_asset_tag_formula_on_device_tabs(self.wb)
        self.assertEqual(tab["G14"].value, "Asset Tag")


if __name__ == "__main__":
    unittest.main()
