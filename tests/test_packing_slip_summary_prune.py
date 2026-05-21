"""Regression test for the Individual packing slip Summary-sheet prune.

Bug: when the user picked a subset of devices and chose "Individual" mode in
the Save dialog, the resulting workbook had the unselected device sheets
deleted — but the Summary sheet was still the full source workbook's
Summary, listing every device from the original inventory upload along with
a stale device-count cell. `_prune_summary_to_selected` rewrites the
Summary to match the surviving device sheets.
"""

import unittest

import openpyxl

from gui.packing_slip_frame import PackingSlipFrame


def _build_workbook_with_summary(devices, asset_tags=None):
    """Build a minimal packing-slip workbook with a Summary sheet + one
    sheet per device, mirroring the uniform layout produced by
    ``build_packing_slip_workbook``: data starts at row 10 with
    B=# | C=IP | D=Device Name (hyperlink) | E=Asset Tag.
    """
    asset_tags = asset_tags or {}
    wb = openpyxl.Workbook()
    summary = wb.active
    summary.title = "Summary"

    # Header block — Customer / Project labels at row 5, values at row 7,
    # Device Count at F5/F7. Column headers at row 9.
    summary["B5"] = "Customer"
    summary["D5"] = "Project"
    summary["F5"] = "Device Count"
    summary["B7"] = "Acme"
    summary["D7"] = "Proj"
    summary["F7"] = len(devices)
    summary["B9"] = "#"
    summary["C9"] = "IP Address"
    summary["D9"] = "Device Name"
    summary["E9"] = "Asset Tag"

    for idx, (ip, device_name) in enumerate(devices):
        wb.create_sheet(device_name)
        row = 10 + idx
        summary[f"B{row}"] = idx + 1
        summary[f"C{row}"] = ip
        cell = summary[f"D{row}"]
        cell.value = device_name
        cell.hyperlink = f"#'{device_name}'!A1"
        cell.style = "Hyperlink"
        if device_name in asset_tags:
            summary[f"E{row}"] = asset_tags[device_name]
    return wb


class TestPruneSummary(unittest.TestCase):

    def test_prune_drops_unselected_devices_and_updates_count(self):
        devices = [
            ("10.0.0.1", "Dev_A"),
            ("10.0.0.2", "Dev_B"),
            ("10.0.0.3", "Dev_C"),
            ("10.0.0.4", "Dev_D"),
        ]
        wb = _build_workbook_with_summary(devices)
        selected = {"Dev_A", "Dev_C"}
        for name in list(wb.sheetnames):
            if name == "Summary" or name in selected:
                continue
            del wb[name]

        PackingSlipFrame._prune_summary_to_selected(wb, ["Summary"], selected)

        ws = wb["Summary"]
        kept_d = [ws.cell(r, 4).value for r in range(10, 15)]
        self.assertEqual(kept_d, ["Dev_A", "Dev_C", None, None, None])
        self.assertEqual(ws["F7"].value, 2)
        self.assertEqual(ws.cell(10, 3).value, "10.0.0.1")
        self.assertEqual(ws.cell(11, 3).value, "10.0.0.3")
        # Sequence renumbered after prune.
        self.assertEqual(ws.cell(10, 2).value, 1)
        self.assertEqual(ws.cell(11, 2).value, 2)
        d_link = ws.cell(10, 4).hyperlink
        self.assertIsNotNone(d_link)
        # openpyxl normalizes internal links to `.location` only after a
        # save/reload cycle; freshly assigned ones live in `.target`.
        self.assertIn("Dev_A", (d_link.location or d_link.target or ""))

    def test_prune_keeps_all_when_all_selected(self):
        devices = [("10.0.0.1", "Dev_A"), ("10.0.0.2", "Dev_B")]
        wb = _build_workbook_with_summary(devices)
        PackingSlipFrame._prune_summary_to_selected(wb, ["Summary"], {"Dev_A", "Dev_B"})
        ws = wb["Summary"]
        self.assertEqual(ws["F7"].value, 2)
        self.assertEqual(ws.cell(10, 4).value, "Dev_A")
        self.assertEqual(ws.cell(11, 4).value, "Dev_B")

    def test_prune_empties_when_nothing_selected(self):
        devices = [("10.0.0.1", "Dev_A"), ("10.0.0.2", "Dev_B")]
        wb = _build_workbook_with_summary(devices)
        PackingSlipFrame._prune_summary_to_selected(wb, ["Summary"], set())
        ws = wb["Summary"]
        self.assertEqual(ws["F7"].value, 0)
        self.assertIsNone(ws.cell(10, 4).value)

    def test_prune_preserves_asset_tags_for_kept_rows(self):
        devices = [
            ("10.0.0.1", "Dev_A"),
            ("10.0.0.2", "Dev_B"),
            ("10.0.0.3", "Dev_C"),
        ]
        tags = {"Dev_A": "AT-100", "Dev_B": "AT-200", "Dev_C": "AT-300"}
        wb = _build_workbook_with_summary(devices, asset_tags=tags)
        PackingSlipFrame._prune_summary_to_selected(wb, ["Summary"], {"Dev_A", "Dev_C"})
        ws = wb["Summary"]
        # Dev_A keeps its tag at the new row 10, Dev_C's tag moves up to row 11.
        self.assertEqual(ws.cell(10, 4).value, "Dev_A")
        self.assertEqual(ws.cell(10, 5).value, "AT-100")
        self.assertEqual(ws.cell(11, 4).value, "Dev_C")
        self.assertEqual(ws.cell(11, 5).value, "AT-300")
        # Dev_B's tag is gone along with Dev_B.
        self.assertIsNone(ws.cell(12, 5).value)


class TestExtractSheetFromLink(unittest.TestCase):

    def test_internal_link_with_fragment(self):
        self.assertEqual(
            PackingSlipFrame._extract_sheet_from_link("#'My Sheet'!A1"),
            "My Sheet",
        )

    def test_internal_link_without_fragment(self):
        self.assertEqual(
            PackingSlipFrame._extract_sheet_from_link("'Sheet1'!B7"),
            "Sheet1",
        )

    def test_unquoted_sheet_name(self):
        self.assertEqual(
            PackingSlipFrame._extract_sheet_from_link("Sheet1!A1"),
            "Sheet1",
        )

    def test_empty_returns_none(self):
        self.assertIsNone(PackingSlipFrame._extract_sheet_from_link(""))
        self.assertIsNone(PackingSlipFrame._extract_sheet_from_link(None))


if __name__ == "__main__":
    unittest.main()
