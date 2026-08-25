"""Tests for Qt-facing packing-slip generation services."""
from unittest.mock import MagicMock

import openpyxl
import pandas as pd

from services.packing_slip_service import PackingSlipRequest, PackingSlipSource, _build_consolidated_workbook, inspect_packing_slip_source, run_packing_slip_generation


def test_inspect_inventory_extracts_summary_metadata(tmp_path) -> None:
    path = tmp_path / "inventory.xlsx"
    workbook = openpyxl.Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["B7"] = "Customer A"
    summary["D7"] = "Project B"
    device = workbook.create_sheet("DEVICE-1")
    device["B7"] = "Customer PO:"
    device["C5"] = "Customer A"
    device["C6"] = "Project B"
    device["C7"] = "PO-1"
    device["D7"] = "SO-2"
    workbook.save(path)
    workbook.close()

    source = inspect_packing_slip_source(path)

    assert source.customer == "Customer A"
    assert source.project == "Project B"
    assert source.purchase_order == "PO-1"
    assert source.sales_order == "SO-2"
    assert source.multisheet


def test_inspect_counts_only_structural_device_sheets(tmp_path) -> None:
    path = tmp_path / "individual.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "Summary"
    for index in range(20):
        sheet = workbook.create_sheet(f"DEVICE-{index + 1}")
        sheet["B7"] = "Device ID:"
        sheet["C7"] = sheet.title
        sheet["D14"] = "Part Number"
        sheet["E14"] = "Serial Number"
    bom = workbook.create_sheet("BOM")
    bom["B7"] = "Device ID:"
    workbook.save(path)
    workbook.close()

    source = inspect_packing_slip_source(path)

    assert source.device_count == 20


def test_consolidated_output_excludes_summary_and_bom(tmp_path) -> None:
    source_path = tmp_path / "individual.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "Summary"
    device = workbook.create_sheet("DEVICE-1")
    device["B7"] = "Device ID:"
    device["D14"] = "Part Number"
    device["E14"] = "Serial Number"
    device["C15"] = "PO-1"
    device["D15"] = "PART-1"
    device["E15"] = "SERIAL-1"
    device["F15"] = "Description"
    bom = workbook.create_sheet("BOM")
    bom["B7"] = "Device ID:"
    bom["D15"] = "SHOULD-NOT-APPEAR"
    workbook.save(source_path)
    workbook.close()
    output = tmp_path / "consolidated.xlsx"
    builder = MagicMock()
    builder._sanitize_cell.side_effect = lambda value: value
    builder.packing_slip_template = ""

    _build_consolidated_workbook(source_path, output, "Cust", "Proj", builder)

    result = openpyxl.load_workbook(output, data_only=True)
    sheet = result.active
    assert sheet["B15"].value == "DEVICE-1"
    assert sheet["D15"].value == "PART-1"
    assert all(
        cell.value != "SHOULD-NOT-APPEAR"
        for row in sheet.iter_rows()
        for cell in row
    )
    result.close()


def test_generation_delegates_processed_devices_to_builder(tmp_path) -> None:
    source_path = tmp_path / "source.csv"
    pd.DataFrame([{"System Name": "DEV-1", "Part Number": "P-1"}]).to_csv(source_path, index=False)
    source = PackingSlipSource(source_path, "Cust", "Proj", "PO", "SO", 1, False)
    request = PackingSlipRequest(source, tmp_path / "output", "Cust", "Proj", "PO", "SO")
    builder = MagicMock()
    builder.build_unified_packing_slip_workbook.return_value = str(tmp_path / "output" / "packing.xlsx")

    output = run_packing_slip_generation(request, builder=builder)

    assert output.name == "packing.xlsx"
    builder.build_unified_packing_slip_workbook.assert_called_once()


def test_generation_requires_customer_and_project(tmp_path) -> None:
    source = PackingSlipSource(tmp_path / "source.csv", "", "", "TBD", "TBD", 0, False)
    request = PackingSlipRequest(source, tmp_path, "", "")
    try:
        run_packing_slip_generation(request, builder=MagicMock())
    except ValueError as exc:
        assert "Customer and Project" in str(exc)
    else:
        raise AssertionError("missing metadata should fail")
